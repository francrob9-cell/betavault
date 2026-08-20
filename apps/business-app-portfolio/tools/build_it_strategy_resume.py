from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(r"C:\Users\franc\OneDrive\Desktop\Resumes 2025")
OUT_DOCX = OUT_DIR / "RAF_Resume_IT_Strategic_Leadership.docx"

INK = RGBColor(28, 38, 52)
BLUE = RGBColor(22, 84, 132)
MUTED = RGBColor(88, 98, 110)
LINE = "D7DEE8"
LIGHT = "EEF4F8"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_font(run, size=9.5, bold=False, color=INK, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(paragraph, text, size=9.5, bold=False, color=INK):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def set_paragraph(paragraph, before=0, after=3, line=1.03):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def section_title(doc, title):
    p = doc.add_paragraph()
    set_paragraph(p, before=7, after=3, line=1)
    add_text(p, title.upper(), size=9.5, bold=True, color=BLUE)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), LINE)
    borders.append(bottom)
    pPr.append(borders)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph(p, before=0, after=2, line=1.03)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    add_text(p, text, size=8.8, color=INK)
    return p


def add_role(doc, title, org, location, dates, bullets):
    p = doc.add_paragraph()
    set_paragraph(p, before=3, after=1, line=1)
    add_text(p, title.upper(), size=9.5, bold=True, color=INK)
    add_text(p, " | ", size=9.5, color=MUTED)
    add_text(p, org, size=9.2, bold=True, color=BLUE)
    add_text(p, " | ", size=9.2, color=MUTED)
    add_text(p, location, size=9.2, color=MUTED)
    add_text(p, " | ", size=9.2, color=MUTED)
    add_text(p, dates, size=9.2, bold=True, color=MUTED)
    for item in bullets:
        bullet(doc, item)


def add_skill_grid(doc, groups):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.allow_autofit = False
    for i, (label, detail) in enumerate(groups):
        if i % 2 == 0:
            row = table.add_row()
            cells = row.cells
        else:
            cells = table.rows[-1].cells
        cell = cells[i % 2]
        cell.width = Inches(3.15)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_border(cell)
        set_cell_shading(cell, "FFFFFF")
        p = cell.paragraphs[0]
        set_paragraph(p, before=0, after=1, line=1)
        add_text(p, label, size=8.7, bold=True, color=BLUE)
        p2 = cell.add_paragraph()
        set_paragraph(p2, before=0, after=2, line=1.02)
        add_text(p2, detail, size=8.2, color=INK)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(title, before=0, after=0, line=1)
    add_text(title, "ROBERT FRANC-KLEIN", size=20, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(subtitle, before=0, after=3, line=1)
    add_text(subtitle, "STRATEGIC IT LEADER | DIGITAL TRANSFORMATION | ENTERPRISE SYSTEMS", size=9.5, bold=True, color=BLUE)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(contact, before=0, after=6, line=1)
    add_text(contact, "(719) 314-7841  |  rafklein0@gmail.com  |  https://rafklein.com  |  PMP, CSM, CAPM, Lean Six Sigma Green Belt", size=8.6, color=MUTED)

    section_title(doc, "Executive Summary")
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=4, line=1.05)
    add_text(
        p,
        "Strategic IT and operations leader with 10+ years guiding enterprise modernization, portfolio delivery, process redesign, and technology adoption across public-sector, healthcare, higher education, and nonprofit environments. Known for translating executive priorities into governed roadmaps, aligning cross-functional teams, managing vendors and budgets, and building practical systems that improve visibility, compliance, and decision-making. Bilingual English/Spanish.",
        size=9.2,
        color=INK,
    )

    section_title(doc, "Strategic Leadership Strengths")
    strengths = [
        ("IT Strategy & Governance", "Roadmaps, operating models, portfolio intake, executive alignment, board and leadership reporting."),
        ("Enterprise Systems", "Workday, Oracle ERP/FIN/HCM, Microsoft Dynamics, NetSuite, Jira, Smartsheet, Google Workspace, Microsoft 365."),
        ("Program Delivery", "Scope, schedule, budget, risk, dependency management, vendor contracts, implementation planning, adoption."),
        ("Process Modernization", "Finance, HR, payroll, procurement, compliance, grants, reporting, and operational workflows."),
        ("Security & Compliance", "Access controls, permission audits, regulated data workflows, audit readiness, controls, documentation."),
        ("Change Leadership", "Training, stakeholder engagement, communication plans, SOPs, executive buy-in, user enablement."),
    ]
    add_skill_grid(doc, strengths)

    section_title(doc, "Professional Experience")
    add_role(
        doc,
        "Senior Strategic Engagement Manager",
        "Roaring Fork Valley Transportation District",
        "Colorado",
        "2021 - Current",
        [
            "Led cross-functional process re-engineering across Finance, HR, and Operations, translating executive priorities into a unified enterprise systems roadmap for a public-sector transportation agency.",
            "Managed complex scopes, budgets, vendor contracts, invoicing, compliance requirements, dependencies, and risk controls across operational technology initiatives.",
            "Built structured training, documentation, adoption plans, and leadership-facing reporting to improve system use, decision visibility, and accountability.",
            "Partnered with executive sponsors and department leaders to prioritize modernization work, resolve blockers, and sequence practical improvements across the organization.",
        ],
    )

    add_role(
        doc,
        "Strategic Initiative Program Manager",
        "Colorado State HCPF DBA The Resource Exchange",
        "Colorado Springs, CO",
        "2017 - 2021",
        [
            "Directed modernization initiatives supporting 30,000+ users in a regulated healthcare environment, improving audit readiness, operational consistency, and cross-department coordination.",
            "Managed program budgets, vendor relationships, contracts, grant tracking, invoicing, reporting, and executive updates for strategic technology and operations initiatives.",
            "Coordinated multi-system integrations with external vendors and internal teams, balancing compliance, security, user needs, and board-level requirements.",
            "Developed progress monitoring tools, gap assessments, and governance routines that helped leaders evaluate priorities and sustain implementation momentum.",
        ],
    )

    add_role(
        doc,
        "Digital Transformation Project Lead",
        "MSU Denver",
        "Denver, CO",
        "2010 - 2017",
        [
            "Led modernization work across Payroll, HR, Finance, and administrative operations, managing full project lifecycles from process analysis through adoption.",
            "Created centralized documentation, onboarding/offboarding workflows, training resources, and policy structures to strengthen operational consistency and system knowledge.",
            "Managed vendor relationships, contracts, service delivery expectations, and technology-enabled operational improvements across business units.",
            "Supported asset management, facilities-related technology, staff enablement, and infrastructure planning for core institutional operations.",
        ],
    )

    add_role(
        doc,
        "Project Coordinator",
        "University of Maryland",
        "Baltimore, MD",
        "2008 - 2010",
        [
            "Supported HR, Finance, Payroll, front-office operations, and university service delivery through coordinated administrative systems and process controls.",
            "Managed vendor deliverables, client approvals, scheduling, and deployment support to keep operational projects moving on time and aligned with stakeholder expectations.",
        ],
    )

    section_title(doc, "Education & Certifications")
    education = doc.add_paragraph()
    set_paragraph(education, before=0, after=2, line=1.03)
    add_text(education, "University of Denver", size=8.9, bold=True, color=BLUE)
    add_text(education, " - Business Project Management & IT Business Analysis, Masters-level study", size=8.9, color=INK)
    education2 = doc.add_paragraph()
    set_paragraph(education2, before=0, after=2, line=1.03)
    add_text(education2, "New Jersey Institute of Technology", size=8.9, bold=True, color=BLUE)
    add_text(education2, " - BS, Information Technology: Database Administration & Middleware Development", size=8.9, color=INK)
    certs = doc.add_paragraph()
    set_paragraph(certs, before=0, after=2, line=1.03)
    add_text(certs, "Certifications: ", size=8.9, bold=True, color=BLUE)
    add_text(certs, "Project Management Professional (PMP), Certified ScrumMaster (CSM), CAPM, Lean Six Sigma Green Belt, CDTO-C", size=8.9, color=INK)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(footer, before=0, after=0, line=1)
        add_text(footer, "Robert Franc-Klein | Strategic IT Leadership", size=7.8, color=MUTED)

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
